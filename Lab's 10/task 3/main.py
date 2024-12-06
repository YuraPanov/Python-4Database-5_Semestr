from docx import Document

document = Document("Task1.docx")
table = document.tables[0]
data = {}

for row in range(1, 4):
    stat_name = table.cell(row, 0).text[1:]
    stat = table.cell(row, 2).text[1:]

    data[stat_name] = stat

print("Данные по ATmega328:")
print(data)
